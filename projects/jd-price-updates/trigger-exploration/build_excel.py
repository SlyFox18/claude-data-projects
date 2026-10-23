"""Excel workbooks for digging into the JD price analysis (same data as the dashboard).
Usage: python build_excel.py prices.duckdb <out_dir>
Writes five workbooks. Derived columns are Excel formulas, so they recalculate if a value or an input is changed;
open them in Excel (or run recalc_excel.ps1) to compute. Inputs a reader may change are yellow."""
import os, sys
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from margin import LABELS, PM, WAVE_NOTES, bucket, load_frame  # noqa: E402

FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FT = Font(name=FONT, size=14, bold=True)
FIN = Font(name=FONT, size=10, color="0000FF")
YELLOW = PatternFill("solid", fgColor="FFFF00")
WRAP = Alignment(wrap_text=True, vertical="top")
FMT = {"money": '$#,##0.00;($#,##0.00);"-"', "money0": '$#,##0;($#,##0);"-"', "pct": '0.0%;-0.0%;"-"',
       "int": '#,##0;(#,##0);"-"', "date": "yyyy-mm-dd", "num": "#,##0.0", "text": "@"}


def new_wb():
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    return wb


def readme(ws, title, snap, lines):
    ws.title = "Read Me"
    ws["A1"] = title; ws["A1"].font = FT
    ws["A2"] = f"JD price data from DP Prod (Fact_PartPriceChange + dim_Parts), changes through {snap}. Snapshot; it does not update."
    ws["A2"].font = F
    r = 4
    for item in lines:
        if isinstance(item, tuple):
            ws.cell(r, 1, item[0]).font = FB
            ws.cell(r, 2, item[1]).font = F
            ws.cell(r, 2).alignment = WRAP
        else:
            ws.cell(r, 1, item).font = FB if item and not item.startswith(" ") else F
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110


def table_sheet(wb, name, cols, rows, widths=None, freeze="B2", style="TableStyleMedium2"):
    """cols: list of (header, kind, key_or_formula). key = DataFrame column; formula = str with {r} for the row.
    Writes the rows as an Excel table (filter buttons, banding)."""
    ws = wb.create_sheet(name)
    for c, (h, kind, _) in enumerate(cols, 1):
        cell = ws.cell(1, c, h); cell.font = FB; cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    recs = rows.to_dict("records")
    for i, rec in enumerate(recs, 2):
        for c, (h, kind, src) in enumerate(cols, 1):
            if isinstance(src, str) and src.startswith("="):
                v = src.format(r=i)
            else:
                v = rec.get(src)
                if isinstance(v, (bool, np.bool_)):
                    v = "Yes" if v else "No"
                elif isinstance(v, float) and np.isnan(v):
                    v = None
                elif isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                elif isinstance(v, (np.integer,)):
                    v = int(v)
                elif isinstance(v, (np.floating,)):
                    v = float(v)
            cell = ws.cell(i, c, v)
            cell.font = F
            if kind in FMT:
                cell.number_format = FMT[kind]
    n = len(recs) + 1
    ref = f"A1:{get_column_letter(len(cols))}{max(n, 2)}"
    tname = "".join(ch for ch in name if ch.isalnum())
    t = Table(displayName=tname, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 42
    for c, (h, kind, _) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(h, 34 if h == "Description" else 13)
    return ws


def col(cols, header):
    return get_column_letter([h for h, _, _ in cols].index(header) + 1)


def main(db, out):
    os.makedirs(out, exist_ok=True)
    df, WAVES, _ = load_frame(db)
    snap = str(df.EffectiveDate.max().date())
    df["descr"] = df.descr.fillna("")
    df["cc"] = df.cc.fillna("")
    df["bucket_old"] = [LABELS[b] for b in df.ob]
    df["bucket_new"] = [LABELS[b] for b in df.nb]
    last_idx = df.groupby("PartNumber").EffectiveDate.transform("max") == df.EffectiveDate
    df["is_current"] = np.where(last_idx, "Yes", "No")
    sold = df[df.sells]
    base_cols = [("Part Number", "text", "PartNumber"), ("Description", "text", "descr"), ("Dealer Group", "text", "grp"),
                 ("Commodity Code", "text", "cc")]
    prices = [("Old List", "money", "o"), ("New List", "money", "n"), ("Old Cost", "money", "oc"), ("New Cost", "money", "nc"),
              ("Old Sell", "money", "os"), ("New Sell", "money", "ns")]

    # ===================== 1. JD discount cuts =====================
    wb = new_wb()
    readme(wb.active, "JD discount cuts", snap, [
        "What this is",
        ("", "Every price update since Aug 2017 where JD raised our dealer cost by at least 1 point more of list price than before "
             "(cost ÷ list went up by 0.01 or more), on parts we sold in the last 12 months. Our sell price is list plus the Price Matrix "
             "markup, so when JD cuts our discount, margin % drops. A few dozen parts carry most of the dollars."),
        "",
        "How to use it",
        ("", "Watch List is sorted by Yearly Margin $ Lost. 'Still Current Price' = Yes means no later JD update has replaced this one, so "
             "the lower margin is what we are living with today. Filter on it, then review the sell price on the top parts."),
        ("", "Summary counts the cuts and dollars per change date with formulas over the Watch List."),
        "",
        "Columns",
        ("Cost ÷ List Before / After", "Our JD dealer cost as a share of JD list price. A higher number means a smaller discount."),
        ("Margin % Before / After", "(sell − cost) ÷ sell, before and after the update."),
        ("Yearly Margin $ Lost", "(margin % before × new sell − new margin $) × 12-month units. What keeping the old margin % would "
                                 "have made, minus what the new prices make, at last year's volume. An estimate, not booked money."),
        ("12-Mo Units / On Hand", "From dim_Parts, company-wide, as of the snapshot. They describe today, not the date of the change."),
        ("Oil / Lube", "Yes for oil, lube, grease and coolant dealer groups. Their cost in JD's price file swings between about 67% and 87% "
                       "of list and is reset outside the file, so treat these rows as 'check first', not as confirmed losses."),
    ])
    rows = sold[sold.disc_cut].copy()
    rows = rows.assign(_lost=rows.lost_year).sort_values("_lost", ascending=False)
    cols = base_cols + [("Change Date", "date", "EffectiveDate")] + prices
    cols += [("Cost ÷ List Before", "pct", "=IFERROR({oc}{{r}}/{ol}{{r}},\"\")"), ("Cost ÷ List After", "pct", "=IFERROR({nc}{{r}}/{nl}{{r}},\"\")"),
             ("Margin % Before", "pct", "=IFERROR(({os}{{r}}-{oc}{{r}})/{os}{{r}},\"\")"),
             ("Margin % After", "pct", "=IFERROR(({ns}{{r}}-{nc}{{r}})/{ns}{{r}},\"\")"),
             ("12-Mo Units", "int", "demand"), ("On Hand", "int", "qoh"),
             ("Yearly Margin $ Lost", "money0", "=IFERROR(({mb}{{r}}*{ns}{{r}}-({ns}{{r}}-{nc}{{r}}))*{u}{{r}},0)"),
             ("Still Current Price", "text", "is_current"), ("Oil / Lube", "text", "oil")]
    L = {k: col(cols, h) for k, h in [("ol", "Old List"), ("nl", "New List"), ("oc", "Old Cost"), ("nc", "New Cost"),
                                       ("os", "Old Sell"), ("ns", "New Sell"), ("mb", "Margin % Before"), ("u", "12-Mo Units")]}
    cols = [(h, k, s.format(**L) if isinstance(s, str) and s.startswith("=") else s) for h, k, s in cols]
    table_sheet(wb, "Watch List", cols, rows, freeze="B2")
    s = wb.create_sheet("Summary", 1)
    s["A1"] = "Discount cuts by change date (parts we sell)"; s["A1"].font = FB
    for c, h in enumerate(["Change Date", "Parts", "Yearly Margin $ Lost", "Still current", "$ Lost excl. Oil / Lube"], 1):
        s.cell(3, c, h).font = FB
    dcol, lcol, ccol = col(cols, "Change Date"), col(cols, "Yearly Margin $ Lost"), col(cols, "Still Current Price")
    ocol = col(cols, "Oil / Lube")
    dates = sorted(rows.EffectiveDate.dt.normalize().unique(), reverse=True)
    dates = [d for d in dates if (rows.EffectiveDate == d).sum() >= 5]
    for i, d in enumerate(dates, 4):
        s.cell(i, 1, pd.Timestamp(d).to_pydatetime()).number_format = FMT["date"]
        s.cell(i, 2, f"=COUNTIFS('Watch List'!{dcol}:{dcol},A{i})").number_format = FMT["int"]
        s.cell(i, 3, f"=SUMIFS('Watch List'!{lcol}:{lcol},'Watch List'!{dcol}:{dcol},A{i})").number_format = FMT["money0"]
        s.cell(i, 4, f"=COUNTIFS('Watch List'!{dcol}:{dcol},A{i},'Watch List'!{ccol}:{ccol},\"Yes\")").number_format = FMT["int"]
        s.cell(i, 5, f"=SUMIFS('Watch List'!{lcol}:{lcol},'Watch List'!{dcol}:{dcol},A{i},'Watch List'!{ocol}:{ocol},\"No\")").number_format = FMT["money0"]
        for c in range(1, 6):
            s.cell(i, c).font = F
    s.cell(len(dates) + 5, 1, "Dates with fewer than 5 cuts are left out of this table; they are all on the Watch List.").font = F
    s.cell(len(dates) + 6, 1, "Oil / lube: JD's price file shows the cost on these parts jumping to ~87% of list, then something outside "
                              "the file resets it to ~67% before the next update. Their 'cuts' are probably not real losses; check before acting.").font = F
    for c, w in zip("ABCDE", (14, 10, 20, 14, 22)):
        s.column_dimensions[c].width = w
    wb.save(os.path.join(out, "JD Discount Cuts.xlsx"))
    print("discount cuts", len(rows))

    # ===================== 2. Price Matrix boundary watch =====================
    wb = new_wb()
    readme(wb.active, "Price Matrix boundary watch", snap, [
        "What this is",
        ("", "Parts we sell whose current JD list price is within 10% below the start of the next Price Matrix bucket. When a JD increase "
             "pushes a part over the line, it gets the next bucket's smaller markup, so our sell price rises less than list."),
        "",
        "How to use it",
        ("", "Change the yellow cell on Inputs to the list increase you expect at the next wave. 'Crosses at Next Wave' and 'Yearly Margin $ "
             "at Risk' recalculate. The Price Matrix sheet holds the buckets and markups (from PRICE MATRIX 7-9-25.csv); edit a markup there "
             "to see the effect of smoothing a boundary."),
        "",
        "Columns",
        ("Markup Now / After Wave", "Price Matrix markup over list for the bucket the list price falls in, looked up on the Price Matrix sheet."),
        ("Actual Markup Now", "Our sell price ÷ list − 1 today. It matches the matrix in the lower buckets; above $500 our sell prices run "
                              "below the matrix markups, so treat dollars at risk there as an upper bound."),
        ("Yearly Margin $ at Risk", "If the part crosses: list after the wave × (markup now − markup after) × 12-month units."),
    ])
    inp = wb.create_sheet("Inputs")
    inp["A1"] = "Next wave list increase"; inp["A1"].font = FB
    inp["B1"] = 0.03; inp["B1"].font = FIN; inp["B1"].fill = YELLOW; inp["B1"].number_format = FMT["pct"]
    inp["C1"] = "Yellow = change me. 3% is the median JD list increase at the June 2026 wave."; inp["C1"].font = F
    inp.column_dimensions["A"].width = 26; inp.column_dimensions["C"].width = 70
    pm = wb.create_sheet("Price Matrix")
    for c, h in enumerate(["Bucket", "From", "To", "Markup"], 1):
        pm.cell(1, c, h).font = FB
    for i, (lo, hi, mk) in enumerate(PM, 2):
        pm.cell(i, 1, LABELS[i - 2]).font = F
        pm.cell(i, 2, lo).number_format = FMT["money"]
        pm.cell(i, 3, None if hi > 1e9 else hi).number_format = FMT["money"]
        pm.cell(i, 4, mk / 100).number_format = '0.00%'
        for c in (2, 3, 4):
            pm.cell(i, c).font = FIN
    pm.cell(len(PM) + 3, 1, "Source: OneDrive Power BI - Vis/Data/PRICE MATRIX 7-9-25.csv, franchise D (same at every branch). Blue = input.").font = F
    for c, w in zip("ABCD", (18, 12, 12, 10)):
        pm.column_dimensions[c].width = w
    n_pm = len(PM) + 1
    rngA, rngB, rngD = f"'Price Matrix'!$A$2:$A${n_pm}", f"'Price Matrix'!$B$2:$B${n_pm}", f"'Price Matrix'!$D$2:$D${n_pm}"
    cur = df.groupby("PartNumber").tail(1)
    cur = cur[cur.sells].copy()
    cur["b"] = bucket(cur.n.to_numpy())
    los = np.array([lo for lo, _, _ in PM] + [np.inf])
    cur["next_from"] = los[cur.b + 1]
    cur = cur[(cur.b < len(PM) - 1) & (cur.n * 1.10 >= cur.next_from)].copy()
    mks = np.array([mk for _, _, mk in PM]) / 100
    after = bucket((cur.n * 1.03).to_numpy())
    cur["_risk"] = np.where(after > cur.b, cur.n * 1.03 * (mks[cur.b] - mks[after]) * cur.demand, 0)
    cur["_dist"] = cur.next_from / cur.n - 1
    cur = cur.sort_values(["_risk", "_dist"], ascending=[False, True])
    cols = base_cols + [("Current List", "money", "n"), ("Current Sell", "money", "ns"), ("Current Cost", "money", "nc"),
                        ("Last JD Change", "date", "EffectiveDate"),
                        ("Bucket", "text", "=INDEX(" + rngA + ",MATCH(ROUND({ln}{{r}},2)," + rngB + ",1))"),
                        ("Markup Now", "pct", "=INDEX(" + rngD + ",MATCH(ROUND({ln}{{r}},2)," + rngB + ",1))"),
                        ("Actual Markup Now", "pct", "=IFERROR({sl}{{r}}/{ln}{{r}}-1,\"\")"),
                        ("Next Bucket Starts", "money", "=INDEX(" + rngB + ",MATCH(ROUND({ln}{{r}},2)," + rngB + ",1)+1)"),
                        ("% Below Next Bucket", "pct", "={nb}{{r}}/{ln}{{r}}-1"),
                        ("List After Next Wave", "money", "=ROUND({ln}{{r}}*(1+Inputs!$B$1),2)"),
                        ("Crosses at Next Wave", "text", "=IF({la}{{r}}>={nb}{{r}},\"Yes\",\"No\")"),
                        ("Markup After Wave", "pct", "=INDEX(" + rngD + ",MATCH({la}{{r}}," + rngB + ",1))"),
                        ("12-Mo Units", "int", "demand"),
                        ("Yearly Margin $ at Risk", "money0", "=IF({cr}{{r}}=\"Yes\",{la}{{r}}*({mn}{{r}}-{ma}{{r}})*{u}{{r}},0)")]
    L = {k: col(cols, h) for k, h in [("ln", "Current List"), ("sl", "Current Sell"), ("nb", "Next Bucket Starts"),
                                       ("la", "List After Next Wave"), ("cr", "Crosses at Next Wave"), ("mn", "Markup Now"),
                                       ("ma", "Markup After Wave"), ("u", "12-Mo Units")]}
    cols = [(h, k, s.format(**L) if isinstance(s, str) and s.startswith("=") else s) for h, k, s in cols]
    table_sheet(wb, "Watch List", cols, cur)
    sm = wb.create_sheet("Summary", 1)
    risk, bcol, crc = col(cols, "Yearly Margin $ at Risk"), col(cols, "Bucket"), col(cols, "Crosses at Next Wave")
    sm["A1"] = "At the increase on Inputs"; sm["A1"].font = FB
    sm["A2"] = "Parts that cross"; sm["B2"] = f"=COUNTIF('Watch List'!{crc}:{crc},\"Yes\")"; sm["B2"].number_format = FMT["int"]
    sm["A3"] = "Yearly margin $ at risk"; sm["B3"] = f"=SUM('Watch List'!{risk}:{risk})"; sm["B3"].number_format = FMT["money0"]
    sm["A5"] = "By current bucket"; sm["A5"].font = FB
    for c, h in enumerate(["Bucket", "Parts that cross", "Yearly margin $ at risk"], 1):
        sm.cell(6, c, h).font = FB
    for i, lab in enumerate(LABELS[:-1], 7):
        sm.cell(i, 1, lab)
        sm.cell(i, 2, f"=COUNTIFS('Watch List'!{bcol}:{bcol},A{i},'Watch List'!{crc}:{crc},\"Yes\")").number_format = FMT["int"]
        sm.cell(i, 3, f"=SUMIFS('Watch List'!{risk}:{risk},'Watch List'!{bcol}:{bcol},A{i})").number_format = FMT["money0"]
    for row in sm.iter_rows(min_row=2, max_row=len(LABELS) + 6):
        for c in row:
            if c.font != FB:
                c.font = F
    for c, w in zip("ABC", (24, 16, 22)):
        sm.column_dimensions[c].width = w
    wb.save(os.path.join(out, "Price Matrix Boundary Watch.xlsx"))
    print("boundary watch", len(cur))

    # ===================== 3. Margin changes by wave =====================
    wb = new_wb()
    readme(wb.active, "Margin changes by wave", snap, [
        "What this is",
        ("", "For each JD repricing wave (40,000+ parts changed on one date): how many parts' margin % held, went down or went up, "
             "and the margin dollars given up by the parts that went down. 'Parts We Sell' = parts with sales in the last 12 months."),
        "",
        "Definitions",
        ("Held / Down / Up", "Margin % = (sell − cost) ÷ sell. Down = margin % fell by more than half a point at that update; Up = rose by "
                             "more than half a point; Held = everything else."),
        ("Yearly Margin $ Lost", "For each part whose margin % went down: (old margin % × new sell − new margin $) × 12-month units, summed. "
                                 "It is what keeping the old margin % would have made at last year's volume, minus what the new prices make."),
        ("Yearly Margin $ Added", "For every part we sell that changed at the wave: (new margin $ − old margin $) per unit × 12-month units, "
                                  "summed. It answers: if we sell the same quantities as last year, how much more gross margin per year do "
                                  "the new prices give us? It is an estimate at today's volume, not money booked."),
        ("Cause", "Oil/lube cost swing = a discount cut on an oil, lube, grease or coolant part (probably not a real loss; their cost "
                  "is reset outside JD's price file). JD discount cut = cost ÷ list rose by 0.01+. Moved into a lower-markup bucket = the update crossed a Price Matrix "
                  "boundary and our sell rose less than list. Sell rose less than list = the same without a bucket change. Other = "
                  "anything else (usually a cost change with list unchanged)."),
    ])
    pdn = sold[(sold.outcome == "down") & sold.EffectiveDate.isin(WAVES)].sort_values(["EffectiveDate", "lost_year"], ascending=[True, False])
    pcols = [("Wave", "date", "EffectiveDate")] + base_cols + [("Cause", "text", "cause")] + prices
    pcols += [("Margin % Before", "pct", "=IFERROR(({os}{{r}}-{oc}{{r}})/{os}{{r}},\"\")"),
              ("Margin % After", "pct", "=IFERROR(({ns}{{r}}-{nc}{{r}})/{ns}{{r}},\"\")"),
              ("12-Mo Units", "int", "demand"),
              ("Yearly Margin $ Lost", "money0", "=IFERROR(({mb}{{r}}*{ns}{{r}}-({ns}{{r}}-{nc}{{r}}))*{u}{{r}},0)")]
    L = {k: col(pcols, h) for k, h in [("oc", "Old Cost"), ("nc", "New Cost"), ("os", "Old Sell"), ("ns", "New Sell"),
                                        ("mb", "Margin % Before"), ("u", "12-Mo Units")]}
    pcols = [(h, k, s.format(**L) if isinstance(s, str) and s.startswith("=") else s) for h, k, s in pcols]
    wc, cc_, lc = col(pcols, "Wave"), col(pcols, "Cause"), col(pcols, "Yearly Margin $ Lost")
    wv = df[df.EffectiveDate.isin(WAVES)]
    summ = []
    for d in WAVES:
        s_ = wv[wv.EffectiveDate == d]; ss = s_[s_.sells]
        summ.append({"wave": d, "parts": len(s_), "held": (s_.outcome == "held").sum(), "down": (s_.outcome == "down").sum(),
                     "up": (s_.outcome == "up").sum(), "sold": len(ss), "sold_up": (ss.outcome == "up").sum(),
                     "added": round(float((ss.demand * (ss.nm - ss.om)).sum())), "cost_med": float(np.nanmedian(s_.cost_chg)),
                     "margin_med": float(np.nanmedian(s_.margin_chg)), "note": WAVE_NOTES.get(str(d.date()), "")})
    pl = "'Parts Margin Down'"
    scols = [("Wave", "date", "wave"), ("Parts Changed", "int", "parts"), ("Held", "int", "held"), ("Down", "int", "down"), ("Up", "int", "up"),
             ("% Held", "pct", "=C{r}/B{r}"), ("% Down", "pct", "=D{r}/B{r}"), ("% Up", "pct", "=E{r}/B{r}"),
             ("Parts We Sell", "int", "sold"),
             ("We Sell: Down", "int", f"=COUNTIFS({pl}!${wc}:${wc},A{{r}})"),
             ("We Sell: % Down", "pct", "=J{r}/I{r}"), ("We Sell: Up", "int", "sold_up"),
             ("Yearly Margin $ Lost (down parts)", "money0", f"=SUMIFS({pl}!${lc}:${lc},{pl}!${wc}:${wc},A{{r}})"),
             ("…from JD discount cuts", "money0", f"=SUMIFS({pl}!${lc}:${lc},{pl}!${wc}:${wc},A{{r}},{pl}!${cc_}:${cc_},\"JD discount cut\")"),
             ("…from bucket moves", "money0", f"=SUMIFS({pl}!${lc}:${lc},{pl}!${wc}:${wc},A{{r}},{pl}!${cc_}:${cc_},\"Moved into a lower-markup bucket\")"),
             ("…from oil/lube cost swings", "money0", f"=SUMIFS({pl}!${lc}:${lc},{pl}!${wc}:${wc},A{{r}},{pl}!${cc_}:${cc_},\"Oil/lube cost swing\")"),
             ("Yearly Margin $ Added (whole wave)", "money0", "added"),
             ("Median Cost % Change", "pct", "cost_med"), ("Median Margin $ % Change", "pct", "margin_med"), ("Note", "text", "note")]
    table_sheet(wb, "Wave Summary", scols, pd.DataFrame(summ), widths={"Note": 60}, freeze="B2")
    table_sheet(wb, "Parts Margin Down", pcols, pdn, freeze="C2")
    wb.save(os.path.join(out, "Margin Changes by Wave.xlsx"))
    print("margin down rows", len(pdn))

    # ===================== 4. Buy-ahead candidates =====================
    wb = new_wb()
    june = [w for w in WAVES if w.year >= 2024]
    readme(wb.active, "Buy-ahead candidates", snap, [
        "What this is",
        ("", "The 1,000 parts we sell with the most yearly cost volume (12-month units × current JD dealer cost), with how much JD changed "
             "their cost at the last three June waves. The main JD wave has landed in mid to late June three years running "
             f"({', '.join(str(w.date()) for w in june)})."),
        "",
        "How to use it",
        ("", "Set the three yellow inputs on Inputs. Each row then shows what buying that many months of demand just before the wave would "
             "cost, save and cost to hold. Summary adds it up for the top 100 / 250 / 500 / 1,000 parts."),
        "",
        "Limits",
        ("", "Assumes the wave date is known in advance and that every part rises by the expected %. Real increases vary by part (see the "
             "June columns). No obsolescence, space or cash limits. Demand is the last 12 months, company-wide."),
    ])
    inp = wb.create_sheet("Inputs")
    for r_, (lab, v, f, note) in enumerate([("Months of demand bought early", 2, "0", "Whole months."),
                                            ("Yearly cost to hold stock", 0.10, FMT["pct"], "Cost of money + storage + risk, per year."),
                                            ("Expected cost increase at next wave", 0.03, FMT["pct"], "June 2026 median was 3%; June 2025 1.5%; June 2024 5%.")], 1):
        inp.cell(r_, 1, lab).font = FB
        c = inp.cell(r_, 2, v); c.font = FIN; c.fill = YELLOW; c.number_format = f
        inp.cell(r_, 3, note).font = F
    inp.column_dimensions["A"].width = 36; inp.column_dimensions["C"].width = 60
    last = df.groupby("PartNumber").tail(1)
    cand = last[last.sells & (last.nc > 0)].copy()
    cand["_vol"] = cand.demand * cand.nc
    cand = cand.sort_values("_vol", ascending=False).head(1000).copy()
    cand["rank"] = range(1, len(cand) + 1)
    for w in june:
        m = df[df.EffectiveDate == w].set_index("PartNumber").cost_chg
        cand[f"jun{w.year}"] = cand.PartNumber.map(m)
    cols = [("Rank", "int", "rank")] + base_cols + [("12-Mo Units", "int", "demand"), ("On Hand", "int", "qoh"),
                                                    ("Current Cost", "money", "nc"), ("Current List", "money", "n")]
    cols += [(f"Cost Change Jun {w.year}", "pct", f"jun{w.year}") for w in june]
    cols += [("Yearly Cost Volume", "money0", "={u}{{r}}*{c}{{r}}"), ("Monthly Units", "num", "={u}{{r}}/12"),
             ("Buy Early Qty", "int", "=ROUND({mu}{{r}}*Inputs!$B$1,0)"), ("Buy Early Cost", "money0", "={q}{{r}}*{c}{{r}}"),
             ("Est. Saving", "money0", "={bc}{{r}}*Inputs!$B$3"),
             ("Holding Cost", "money0", "={mu}{{r}}*{c}{{r}}*Inputs!$B$2/12*Inputs!$B$1*(Inputs!$B$1+1)/2"),
             ("Net", "money0", "={sv}{{r}}-{hc}{{r}}")]
    L = {k: col(cols, h) for k, h in [("u", "12-Mo Units"), ("c", "Current Cost"), ("mu", "Monthly Units"), ("q", "Buy Early Qty"),
                                       ("bc", "Buy Early Cost"), ("sv", "Est. Saving"), ("hc", "Holding Cost")]}
    cols = [(h, k, s.format(**L) if isinstance(s, str) and s.startswith("=") else s) for h, k, s in cols]
    table_sheet(wb, "Candidates", cols, cand, freeze="C2")
    sm = wb.create_sheet("Summary", 1)
    rk = col(cols, "Rank")
    for c, h in enumerate(["Top N parts", "Buy Early Cost", "Est. Saving", "Holding Cost", "Net"], 1):
        sm.cell(1, c, h).font = FB
    for i, nn in enumerate((100, 250, 500, 1000), 2):
        sm.cell(i, 1, nn).font = F
        for c, h in enumerate(["Buy Early Cost", "Est. Saving", "Holding Cost", "Net"], 2):
            k = col(cols, h)
            sm.cell(i, c, f"=SUMIFS(Candidates!{k}:{k},Candidates!{rk}:{rk},\"<=\"&A{i})").number_format = FMT["money0"]
            sm.cell(i, c).font = F
    for c, w in zip("ABCDE", (14, 18, 14, 14, 14)):
        sm.column_dimensions[c].width = w
    wb.save(os.path.join(out, "Buy-Ahead Candidates.xlsx"))
    print("buy-ahead", len(cand))

    # ===================== 5. Every change since 2024, parts we sell =====================
    wb = new_wb()
    readme(wb.active, "JD price changes since 2024 — parts we sell", snap, [
        "What this is",
        ("", "Every JD price update since 1 Jan 2024 on parts with sales in the last 12 months, one row per part per update. The raw "
             "material behind the other workbooks, for your own filters and pivots."),
        "",
        "Columns",
        ("List / Cost / Sell % Change", "New ÷ old − 1."),
        ("Margin Outcome", "Held / Down / Up: margin % moved by less or more than half a point."),
        ("Bucket Before / After", "Price Matrix bucket of the old and new list price."),
        ("JD Discount Cut", "Yes when cost ÷ list rose by 0.01 or more."),
        ("12-Mo Units / On Hand", "From dim_Parts as of the snapshot, company-wide."),
    ])
    hist = sold[sold.EffectiveDate >= "2024-01-01"].sort_values(["PartNumber", "EffectiveDate"])
    cols = base_cols + [("Change Date", "date", "EffectiveDate")] + prices
    cols += [("List % Change", "pct", "=IFERROR({nl}{{r}}/{ol}{{r}}-1,\"\")"), ("Cost % Change", "pct", "=IFERROR({nc}{{r}}/{oc}{{r}}-1,\"\")"),
             ("Sell % Change", "pct", "=IFERROR({ns}{{r}}/{os}{{r}}-1,\"\")"),
             ("Margin % Before", "pct", "=IFERROR(({os}{{r}}-{oc}{{r}})/{os}{{r}},\"\")"),
             ("Margin % After", "pct", "=IFERROR(({ns}{{r}}-{nc}{{r}})/{ns}{{r}},\"\")"),
             ("Margin Outcome", "text", "outcome"), ("Bucket Before", "text", "bucket_old"), ("Bucket After", "text", "bucket_new"),
             ("JD Discount Cut", "text", "disc_cut"), ("Oil / Lube", "text", "oil"), ("12-Mo Units", "int", "demand"), ("On Hand", "int", "qoh")]
    L = {k: col(cols, h) for k, h in [("ol", "Old List"), ("nl", "New List"), ("oc", "Old Cost"), ("nc", "New Cost"),
                                       ("os", "Old Sell"), ("ns", "New Sell")]}
    cols = [(h, k, s.format(**L) if isinstance(s, str) and s.startswith("=") else s) for h, k, s in cols]
    table_sheet(wb, "Changes", cols, hist)
    wb.save(os.path.join(out, "JD Price Changes Since 2024 - Parts We Sell.xlsx"))
    print("history", len(hist))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv[1], sys.argv[2])
