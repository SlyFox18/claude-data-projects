"""
COTTON HARVESTER INSPECTION HISTORY (AD HOC) — for Casey
============================================================================
Request (2026-10-21): Casey and the CEO are visiting locations to talk about
inspections. Casey wants the inspection history for:
  Branches:   6 (San Angelo), 8 (Big Spring), 96 (Snyder), 13 (Lamesa)
  Dates:      2025-08-01 to today (work order creation date)
  Machines:   CS770, CS690 (asked as "CCS690"), 7460, CP770, CP690
Delivered as an Excel workbook, not a report.

Source: DP_Presentation (Prod) lakehouse, read directly via DuckDB over
OneLake — the same data the Inspections report reads (Fact_LaborJobSummary
is built from Silver_WkOthSub / WkMechWk / WkRoFile; this pulls those Silver
tables through DP_Presentation's shortcuts so it can add customer, serial
number, machine model and parts, which the Gold table doesn't carry).

WHICH JOBS COUNT AS "AN INSPECTION" HERE
- Any job code in lookup_InspectionJobCodes (the report's own list of 113
  inspection codes), EXCEPT IS-AMS SOFTWARE (a software update, not an
  inspection), PLUS the JD standard "Pre Harvest Inspection, Adjust &
  Calibrate" codes (<model>/9072/150) which aren't in the report's list.
- AND the job is for one of the requested machines, matched EITHER by:
    * a model-specific inspection code (IS-CS770 / IS-CS690 / IS-CP770 /
      IS-CP690 INSPECT, /CS690 INSPECTION, /CS690 VIP INSPECTIO), OR
    * the machine on the work order (Silver_WkVehFl.Model via the work
      order's Registration) being one of the five models.
  The model match matters: there is NO 7460-specific inspection code —
  7460s are inspected under the generic IS-STRIPPER INSPECT (CS690s
  sometimes are too). Filtering on job code alone would miss every 7460.

Output: "Cotton Harvester Inspections - Br 6, 8, 13, 96 (2025-08-01 to <today>).xlsx"
- Summary:          branch x machine model — inspections, completed vs open,
                    inspection labor, follow-up work, parts, total WO $.
- Inspections:      one row per inspection job, with customer, serial, model,
                    status, hours, $ and the totals for its whole work order.
- All Work on WOs:  every job line on those work orders (the inspection
                    itself + everything found/added — the "hitlist" work).
- Parts on WOs:     part lines invoiced on those work orders (from
                    Fact_WorkOrderParts). Reference detail only: that table is
                    missing some invoices (e.g. Br 8 WO 689171 / invoice
                    1950991 has no lines at all; ~$43K short over the 63 WOs
                    as of 2026-10-21), so $ totals use WkOthSub.InvParts.
- Notes:            scope and definitions for the reader.

Run manually — not part of any scheduled pipeline:
    timeout 600 python build_report.py
============================================================================
"""

import datetime as dt
import re

import duckdb
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WS_ID = "7836042d-adb1-4846-b70d-bd42980054c5"  # DP - Presentation - Prod
LH_ID = "29d9df80-a383-4d40-9807-1e2e6cbff88f"  # DP_Presentation (Prod)
base = f"abfss://{WS_ID}@onelake.dfs.fabric.microsoft.com/{LH_ID}/Tables"

BRANCHES = ["6", "8", "96", "13"]
START_DATE = "2025-08-01"
TODAY = dt.date.today().isoformat()
MODELS = ["7460", "CS690", "CS770", "CP690", "CP770"]
# Machine-file model variants that are the same machine (7460AU, CS690WW, ...).
# Attachments/kits ("7460 HEADER", "CS690 WALKWAY") deliberately don't match.
MODEL_REGEX = r"^(7460|CS690|CS770|CP690|CP770)(AU|U|WW|RU)?$"
MODEL_SPECIFIC_CODES = {
    "IS-CS770 INSPECT": "CS770", "IS-CS690 INSPECT": "CS690",
    "IS-CP770 INSPECT": "CP770", "IS-CP690 INSPECT": "CP690",
    "/CS690 INSPECTION": "CS690", "/CS690 VIP INSPECTIO": "CS690",
}
EXCLUDED_CODES = ["IS-AMS SOFTWARE"]
PRE_HARVEST_CODES = [f"{m}/9072/150" for m in MODELS]
JOB_TYPES = {"R": "Customer", "I": "Internal", "W": "Warranty", "S": "S"}

out_path = f"Cotton Harvester Inspections - Br 6, 8, 13, 96 ({START_DATE} to {TODAY}).xlsx"

con = duckdb.connect()
con.execute("INSTALL delta; LOAD delta; INSTALL azure; LOAD azure;")
con.execute("CREATE SECRET (TYPE azure, PROVIDER credential_chain, CHAIN 'cli');")

for name in ["Silver_WkOthSub", "Silver_WkRoFile", "Silver_WkVehFl", "Silver_WkMechWk",
             "Silver_WkCodeFl", "lookup_InspectionJobCodes", "dim_CustomerList",
             "dim_BranchLocation", "Fact_WorkOrderParts", "dim_Parts"]:
    con.execute(f"CREATE VIEW {name} AS SELECT * FROM delta_scan('{base}/{name}')")

sql_list = lambda xs: ", ".join(f"'{x}'" for x in xs)
con.register("model_codes", pd.DataFrame(
    {"JobCode": list(MODEL_SPECIFIC_CODES), "CodeModel": list(MODEL_SPECIFIC_CODES.values())}))

# ------------------------------------------------------------------
# Work orders at the 4 branches created in the window, with machine + customer
# ------------------------------------------------------------------
con.execute(f"""
    CREATE TEMP TABLE wo AS
    SELECT r.Branch, r.WorkOrder,
           CAST(r.CreatedOn AS DATE) AS WOCreated,
           CAST(r.ClosedDate AS DATE) AS WOClosed,
           CASE WHEN r.IsClosed = 'Y' THEN 'Closed' ELSE 'Open' END AS WOStatus,
           r.Registration, r.AccountNumber, r.Salesperson,
           v.VIN, v.Make, UPPER(TRIM(v.Model)) AS MachineModel, v.Year AS ModelYear
    FROM Silver_WkRoFile r
    LEFT JOIN Silver_WkVehFl v ON v.Registration = r.Registration
    WHERE r.Branch IN ({sql_list(BRANCHES)})
      AND r.CreatedOn >= TIMESTAMP '{START_DATE}'
""")

# Job lines (WkOthSub grain) on those work orders, with punched hours.
con.execute(f"""
    CREATE TEMP TABLE jobs AS
    WITH hours AS (
        SELECT Branch, WorkOrder, JobCode, JobType, SUM(COALESCE(HoursWorked, 0)) AS HoursWorked
        FROM Silver_WkMechWk GROUP BY ALL
    ),
    code_desc AS (
        SELECT TRIM(CODE) AS JobCode,
               MIN(TRIM(REGEXP_REPLACE(SPLIT_PART(DESCRIPTION, CHR(13), 1), '\\s+', ' ', 'g'))) AS JobDescription
        FROM Silver_WkCodeFl GROUP BY 1
    )
    SELECT o.Branch, o.WorkOrder, TRIM(o.JobCode) AS JobCode, o.JobType,
           d.JobDescription,
           CASE o.JobStatus WHEN 'C' THEN 'Complete' WHEN 'O' THEN 'Open' ELSE o.JobStatus END AS JobStatus,
           o.InvoiceNumber, CAST(o.InvoiceDate AS DATE) AS InvoiceDate,
           COALESCE(h.HoursWorked, 0) AS HoursWorked, o.EstHours,
           COALESCE(o.InvLabor, 0) AS InvLabor, COALESCE(o.InvParts, 0) AS InvParts,
           COALESCE(o.EstLabor, 0) AS EstLabor, COALESCE(o.EstParts, 0) AS EstParts,
           (TRIM(o.JobCode) IN (SELECT TRIM(JobCode) FROM lookup_InspectionJobCodes)
               AND TRIM(o.JobCode) NOT IN ({sql_list(EXCLUDED_CODES)}))
             OR TRIM(o.JobCode) IN ({sql_list(PRE_HARVEST_CODES)}) AS IsInspectionJob
    FROM Silver_WkOthSub o
    JOIN wo ON wo.Branch = o.Branch AND wo.WorkOrder = o.WorkOrder
    LEFT JOIN hours h ON h.Branch = o.Branch AND h.WorkOrder = o.WorkOrder
                     AND h.JobCode = o.JobCode AND h.JobType = o.JobType
    LEFT JOIN code_desc d ON d.JobCode = TRIM(o.JobCode)
""")

# Inspection jobs for the requested machines.
con.execute(f"""
    CREATE TEMP TABLE insp AS
    SELECT j.*, wo.MachineModel, mc.CodeModel,
           COALESCE(
               CASE WHEN regexp_matches(wo.MachineModel, '{MODEL_REGEX}')
                    THEN regexp_extract(wo.MachineModel, '{MODEL_REGEX}', 1) END,
               mc.CodeModel,
               regexp_extract(j.JobCode, '^(7460|CS690|CS770|CP690|CP770)/9072/150$', 1)
           ) AS Model,
           CASE WHEN mc.CodeModel IS NOT NULL AND regexp_matches(wo.MachineModel, '{MODEL_REGEX}') THEN 'Job code + machine'
                WHEN mc.CodeModel IS NOT NULL THEN 'Job code'
                ELSE 'Machine model' END AS MatchedOn
    FROM jobs j
    JOIN wo ON wo.Branch = j.Branch AND wo.WorkOrder = j.WorkOrder
    LEFT JOIN model_codes mc ON mc.JobCode = j.JobCode
    WHERE j.IsInspectionJob
      AND (mc.CodeModel IS NOT NULL
           OR regexp_matches(wo.MachineModel, '{MODEL_REGEX}')
           OR regexp_matches(j.JobCode, '^(7460|CS690|CS770|CP690|CP770)/9072/150$'))
""")

# Work-order totals for every work order with a qualifying inspection.
con.execute("""
    CREATE TEMP TABLE wo_totals AS
    SELECT j.Branch, j.WorkOrder,
           COUNT(*) FILTER (WHERE NOT j.IsInspectionJob) AS OtherJobLines,
           SUM(j.InvLabor) FILTER (WHERE j.IsInspectionJob) AS InspLabor,
           SUM(j.InvLabor) FILTER (WHERE NOT j.IsInspectionJob) AS FollowUpLabor,
           SUM(j.InvParts) AS WOParts,
           SUM(j.InvLabor + j.InvParts) AS WOTotal,
           SUM(j.EstLabor + j.EstParts) FILTER (WHERE j.JobStatus = 'Open') AS OpenEstimate,
           SUM(j.HoursWorked) AS WOHours
    FROM jobs j
    WHERE (j.Branch, j.WorkOrder) IN (SELECT Branch, WorkOrder FROM insp)
    GROUP BY ALL
""")

branch_name = """(SELECT ANY_VALUE(Branch) FROM dim_BranchLocation b WHERE b.BranchID = {col})"""

inspections = con.execute(f"""
    SELECT {branch_name.format(col='i.Branch')} AS Branch,
           i.WorkOrder AS "Work Order",
           wo.WOCreated AS "WO Created", wo.WOStatus AS "WO Status", wo.WOClosed AS "WO Closed",
           i.Model AS "Model", wo.MachineModel AS "Machine File Model", wo.ModelYear AS "Model Year",
           wo.VIN AS "Serial / PIN", wo.Registration AS "Equip Registration",
           CAST(CAST(wo.AccountNumber AS BIGINT) AS VARCHAR) AS "Customer #",
           c.CustomerName AS "Customer", c.City AS "Customer City",
           i.JobCode AS "Inspection Job Code", i.JobDescription AS "Job Description",
           CASE i.JobType {' '.join(f"WHEN '{k}' THEN '{v}'" for k, v in JOB_TYPES.items())} ELSE i.JobType END AS "Job Type",
           i.JobStatus AS "Inspection Status",
           i.InvoiceNumber AS "Invoice #", i.InvoiceDate AS "Invoice Date",
           ROUND(i.HoursWorked, 2) AS "Inspection Hours", i.EstHours AS "Est Hours",
           i.InvLabor AS "Inspection Labor $",
           t.OtherJobLines AS "Other Jobs on WO",
           COALESCE(t.FollowUpLabor, 0) AS "Follow-up Labor $",
           COALESCE(t.WOParts, 0) AS "WO Parts $",
           COALESCE(t.WOTotal, 0) AS "WO Total Invoiced $",
           COALESCE(t.OpenEstimate, 0) AS "Open Jobs Estimate $",
           i.MatchedOn AS "Matched On"
    FROM insp i
    JOIN wo ON wo.Branch = i.Branch AND wo.WorkOrder = i.WorkOrder
    JOIN wo_totals t ON t.Branch = i.Branch AND t.WorkOrder = i.WorkOrder
    LEFT JOIN dim_CustomerList c ON c.AccountNumber = CAST(CAST(wo.AccountNumber AS BIGINT) AS VARCHAR)
    ORDER BY i.Branch, i.Model, wo.WOCreated, i.WorkOrder
""").df()

all_work = con.execute(f"""
    SELECT {branch_name.format(col='j.Branch')} AS Branch,
           j.WorkOrder AS "Work Order", wo.WOCreated AS "WO Created",
           (SELECT ANY_VALUE(Model) FROM insp i WHERE i.Branch = j.Branch AND i.WorkOrder = j.WorkOrder) AS "Model",
           c.CustomerName AS "Customer",
           j.JobCode AS "Job Code", j.JobDescription AS "Job Description",
           CASE j.JobType {' '.join(f"WHEN '{k}' THEN '{v}'" for k, v in JOB_TYPES.items())} ELSE j.JobType END AS "Job Type",
           CASE WHEN j.IsInspectionJob THEN 'Inspection' ELSE 'Follow-up / other' END AS "Line",
           j.JobStatus AS "Job Status", j.InvoiceNumber AS "Invoice #", j.InvoiceDate AS "Invoice Date",
           ROUND(j.HoursWorked, 2) AS "Hours", j.EstHours AS "Est Hours",
           j.InvLabor AS "Labor $", j.InvParts AS "Parts $",
           j.EstLabor AS "Est Labor $", j.EstParts AS "Est Parts $"
    FROM jobs j
    JOIN wo ON wo.Branch = j.Branch AND wo.WorkOrder = j.WorkOrder
    LEFT JOIN dim_CustomerList c ON c.AccountNumber = CAST(CAST(wo.AccountNumber AS BIGINT) AS VARCHAR)
    WHERE (j.Branch, j.WorkOrder) IN (SELECT Branch, WorkOrder FROM insp)
    ORDER BY j.Branch, j.WorkOrder, j.IsInspectionJob DESC, j.JobCode
""").df()

# Part lines on the inspection work orders' invoices. Fact_WorkOrderParts'
# own Description is "Inv No. ..." - real part description comes from dim_Parts.
parts = con.execute(f"""
    WITH inv AS (
        SELECT DISTINCT j.Branch, j.WorkOrder, j.InvoiceNumber
        FROM jobs j
        WHERE j.InvoiceNumber IS NOT NULL AND TRIM(j.InvoiceNumber) <> ''
          AND (j.Branch, j.WorkOrder) IN (SELECT Branch, WorkOrder FROM insp)
    ),
    pdesc AS (
        SELECT PartNumber, Franchise, ANY_VALUE(Description) AS Description FROM dim_Parts GROUP BY ALL
    )
    SELECT {branch_name.format(col='inv.Branch')} AS Branch,
           inv.WorkOrder AS "Work Order",
           (SELECT ANY_VALUE(Model) FROM insp i WHERE i.Branch = inv.Branch AND i.WorkOrder = inv.WorkOrder) AS "Model",
           p.InvoiceNumber AS "Invoice #", p.InvoiceDate AS "Invoice Date",
           p.PartNumber AS "Part Number", d.Description AS "Part Description", p.Franchise,
           p.Quantity AS "Qty", p.SaleValue AS "Sale $"
    FROM inv
    JOIN Fact_WorkOrderParts p ON p.BranchCode = inv.Branch AND p.InvoiceNumber = inv.InvoiceNumber
    LEFT JOIN pdesc d ON d.PartNumber = p.PartNumber AND d.Franchise = p.Franchise
    ORDER BY inv.Branch, inv.WorkOrder, p.PartNumber
""").df()

# ------------------------------------------------------------------
# Summary: branch x model, counted at work-order grain (a WO with two
# inspection codes counts once).
# ------------------------------------------------------------------
wo_level = (
    inspections.groupby(["Branch", "Model", "Work Order"], as_index=False)
    .agg(**{
        "WO Status": ("WO Status", "first"),
        "Inspection Labor $": ("Inspection Labor $", "sum"),
        "Follow-up Labor $": ("Follow-up Labor $", "first"),
        "WO Parts $": ("WO Parts $", "first"),
        "WO Total Invoiced $": ("WO Total Invoiced $", "first"),
        "Open Jobs Estimate $": ("Open Jobs Estimate $", "first"),
        "Customer #": ("Customer #", "first"),
        "Other Jobs on WO": ("Other Jobs on WO", "first"),
    })
)


def summarize(g):
    n = g["Work Order"].nunique()
    return pd.Series({
        "Inspection WOs": n,
        "Closed": (g["WO Status"] == "Closed").sum(),
        "Open": (g["WO Status"] == "Open").sum(),
        "Customers": g["Customer #"].nunique(),
        "WOs with Follow-up Work": (g["Other Jobs on WO"] > 0).sum(),
        "Inspection Labor $": g["Inspection Labor $"].sum(),
        "Follow-up Labor $": g["Follow-up Labor $"].sum(),
        "Parts $": g["WO Parts $"].sum(),
        "Total Invoiced $": g["WO Total Invoiced $"].sum(),
        "Avg Invoiced $ per WO": g["WO Total Invoiced $"].sum() / n if n else 0,
        "Open Jobs Estimate $": g["Open Jobs Estimate $"].sum(),
    })


by_model = wo_level.groupby(["Branch", "Model"]).apply(summarize, include_groups=False).reset_index()
branch_tot = wo_level.groupby("Branch").apply(summarize, include_groups=False).reset_index()
branch_tot["Model"] = "All models"
grand = summarize(wo_level).to_frame().T
grand["Branch"], grand["Model"] = "All branches", "All models"
summary = pd.concat([by_model, branch_tot, grand], ignore_index=True)
summary["_o"] = summary["Model"].eq("All models")
summary = summary.sort_values(["Branch", "_o", "Model"]).drop(columns="_o")
summary = pd.concat([summary[summary["Branch"] != "All branches"], summary[summary["Branch"] == "All branches"]])
count_cols = ["Inspection WOs", "Closed", "Open", "Customers", "WOs with Follow-up Work"]
summary[count_cols] = summary[count_cols].astype(int)

notes = pd.DataFrame({"Notes": [
    f"Cotton harvester inspection history — branches 6 San Angelo, 8 Big Spring, 13 Lamesa, 96 Snyder.",
    f"Work orders created {START_DATE} through {TODAY}. Machines: 7460, CS690, CS770, CP690, CP770.",
    "Source: the same data the Inspections report uses (DP_Presentation lakehouse, refreshed nightly).",
    "",
    "What counts as an inspection:",
    "• Any job code the Inspections report treats as an inspection (except IS-AMS SOFTWARE, a software update),",
    "  plus John Deere 'Pre Harvest Inspection, Adjust & Calibrate' codes (<model>/9072/150).",
    "• For one of the five machines — by a model-specific inspection code (e.g. IS-CS690 INSPECT) OR by the machine",
    "  on the work order. 7460s have no inspection code of their own; they are booked as IS-STRIPPER INSPECT,",
    "  so they are found by machine model. 'Matched On' shows which rule found each row.",
    "• 'Model' = the machine on the work order when known, otherwise the model in the inspection code.",
    "",
    "Columns:",
    "• Inspection Labor $ = invoiced labor on the inspection job itself.",
    "• Follow-up Labor $ = invoiced labor on every other job on the same work order (repairs found by the inspection).",
    "• WO Parts $ / WO Total Invoiced $ = parts and labor+parts invoiced on the whole work order.",
    "• Open Jobs Estimate $ = estimated labor+parts on jobs not yet complete (quoted work still pending).",
    "• Job Type: Customer (R), Internal (I), Warranty (W).",
    "• Summary counts each work order once, even if it carries two inspection job codes.",
    "• Negative parts $ on a job are discounts/credits applied on the invoice.",
    "• Parts on WOs is part-number detail for reference. It does not list every invoice line (some invoices",
    "  are missing from that table), so it won't add up to WO Parts $. Use the job-line $ for totals.",
    "• Not included: 7455 strippers (2 inspections at Snyder and Lamesa) — not on the requested model list.",
    "",
    "Tabs: Summary · Inspections (one row per inspection job) · All Work on WOs (every job line on those work orders)",
    "· Parts on WOs (part lines invoiced on those work orders).",
]})

# ------------------------------------------------------------------
# Write + format
# ------------------------------------------------------------------
_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
for df in (inspections, all_work, parts):
    for col in df.select_dtypes(include=["object", "string"]).columns:
        df[col] = df[col].map(lambda v: _ILLEGAL.sub("", v) if isinstance(v, str) else v)

HEADER_FILL = PatternFill("solid", fgColor="367C2B")
TOTAL_FILL = PatternFill("solid", fgColor="E2EFDA")
MONEY = '"$"#,##0.00;("$"#,##0.00);"-"'

with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
    sheets = [("Summary", summary), ("Inspections", inspections), ("All Work on WOs", all_work),
              ("Parts on WOs", parts), ("Notes", notes)]
    for name, df in sheets:
        df.to_excel(writer, sheet_name=name, index=False)
        ws = writer.sheets[name]
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        if name == "Notes":
            ws.column_dimensions["A"].width = 130
            continue
        ws.freeze_panes = "A2" if name == "Summary" else "C2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(df.columns, start=1):
            letter = get_column_letter(idx)
            sample = [len(str(v)) for v in df[col].head(500) if pd.notna(v)]
            ws.column_dimensions[letter].width = min(max([len(col) * 0.9, *sample, 8]) + 2, 45)
            fmt = MONEY if "$" in col else ("yyyy-mm-dd" if col in ("WO Created", "WO Closed", "Invoice Date") else None)
            if fmt:
                for (c,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    c.number_format = fmt
        if name == "Summary":
            for row in ws.iter_rows(min_row=2):
                if row[1].value == "All models":
                    for c in row:
                        c.font = Font(bold=True)
                        c.fill = TOTAL_FILL

print(f"Inspection job rows: {len(inspections)}  (WOs: {inspections['Work Order'].nunique()})")
print(inspections.groupby(["Branch", "Model", "Matched On"]).size().to_string())
print(f"All job lines on those WOs: {len(all_work)}")
print(f"Part lines: {len(parts)}")
print(summary.to_string())
print(f"Saved: {out_path}")
