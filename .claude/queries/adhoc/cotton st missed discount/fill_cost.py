"""
COTTON ST MISSED DISCOUNT - FILL PART COST (AD HOC, ONE-TIME)
============================================================================
Request: workbook of 253 order lines (PART #, QTY, COST, TOTAL COST, ORDER #)
with COST blank. Fill COST only; every other cell is left untouched.

Cost = InMaster.STK_ORDER_PRICE (JD dealer cost; identical to
jdis_Part_Information.Cost), Franchise D, from DP_Staging bronze via DuckDB.

Branch rule (from Ben, 2026-10-08): when branches disagree, use store 1 -
"store 1 is usually right"; stores starting with 9 (92/93/95 here) often carry
stale prices. Fallback if a part isn't at store 1: most common cost across
non-9x stores.

Two part numbers on the original sheet were typos (N125108 -> N125408,
N198828 -> N129828); Brian corrected them in the workbook before this run.

Output: a copy with "- WITH COST" in the name; the source file is not modified.
Run manually - not part of any pipeline.
============================================================================
"""
import shutil
import duckdb
import openpyxl
import pandas as pd

SRC = "COTTON ST MISSED DISCOUNT.xlsx"
OUT = "COTTON ST MISSED DISCOUNT - WITH COST.xlsx"
b = ("abfss://ab15d64d-c7ba-415d-9bcf-7feb1ef9b201@onelake.dfs.fabric.microsoft.com/"
     "876255e0-d462-4697-adc1-4a655f5bb101/Tables")

wb = openpyxl.load_workbook(SRC)
ws = wb.active
assert [c.value for c in ws[1]] == ["PART #", "QTY", "COST", "TOTAL COST", "ORDER #"]
rows = list(range(2, ws.max_row + 1))
pns = sorted({str(ws.cell(r, 1).value).strip() for r in rows if ws.cell(r, 1).value})

con = duckdb.connect()
con.execute("INSTALL delta; LOAD delta; INSTALL azure; LOAD azure;")
con.execute("CREATE SECRET (TYPE azure, PROVIDER credential_chain, CHAIN 'cli');")
con.register("p", pd.DataFrame({"pn": pns}))
costs = con.execute(f"""
    WITH im AS (
        SELECT TRIM(PART_NO) pn, TRIM(BRANCH) br, STK_ORDER_PRICE unit_cost
        FROM delta_scan('{b}/InMaster') i JOIN p ON TRIM(i.PART_NO) = p.pn
        WHERE FRANCHISE = 'D'),
    store1 AS (SELECT pn, MAX(unit_cost) unit_cost FROM im WHERE br = '1' GROUP BY pn),
    common AS (
        SELECT pn, unit_cost FROM (
            SELECT pn, unit_cost, ROW_NUMBER() OVER (PARTITION BY pn ORDER BY COUNT(*) DESC, unit_cost DESC) rn
            FROM im WHERE br NOT LIKE '9%' GROUP BY pn, unit_cost) WHERE rn = 1)
    SELECT p.pn, COALESCE(s.unit_cost, c.unit_cost) unit_cost,
           CASE WHEN s.unit_cost IS NOT NULL THEN 'store 1' WHEN c.unit_cost IS NOT NULL THEN 'most common' END src,
           (SELECT COUNT(DISTINCT unit_cost) FROM im WHERE im.pn = p.pn) nvals
    FROM p LEFT JOIN store1 s USING (pn) LEFT JOIN common c USING (pn)
""").df().set_index("pn")

print(costs[costs["nvals"] != 1].to_string())
missing = costs[costs["unit_cost"].isna()].index.tolist()
print(f"distinct parts {len(pns)}, priced {costs['unit_cost'].notna().sum()}, not found {missing}")

for r in rows:
    pn = str(ws.cell(r, 1).value).strip()
    c = costs.loc[pn, "unit_cost"]
    if pd.notna(c):
        ws.cell(r, 3).value = round(float(c), 2)
wb.save(OUT)
print(f"wrote {OUT}")
